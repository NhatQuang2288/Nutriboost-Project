#!/usr/bin/env python3
"""
Đọc packages/seed/source/vdd-tong-hop.xlsx và sinh packages/seed/src/data/vdd.generated.ts.

Cần: pip install openpyxl
Chạy: python3 scripts/import-vdd-xlsx.py

FILE SINH RA KHÔNG ĐƯỢC SỬA TAY — sửa tệp xlsx rồi chạy lại script này.

Số liệu dinh dưỡng của MÓN lấy nguyên từ bảng VDD (theo khẩu phần ghi ở cột "Khối lượng").
Gram từng nguyên liệu là số ƯỚC TÍNH (do người soạn tệp ghi rõ), không phải số gốc.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "packages/seed/source/vdd-tong-hop.xlsx"
OUT = ROOT / "packages/seed/src/data/vdd.generated.ts"

wb = openpyxl.load_workbook(SRC, data_only=True)
dishes_ws, detail_ws = wb.worksheets[0], wb.worksheets[1]


# Ô trống ở đạm/bột đường nghĩa là bằng 0 trong bảng VDD (dầu ăn không có đạm, thịt luộc không có bột đường).
def num(value):
    return None if value is None or value == "" else float(value)


components = {}
for row in detail_ws.iter_rows(min_row=5, values_only=True):
    if not row[0]:
        continue
    components.setdefault(row[1], []).append({"name": str(row[2]).strip(), "grams": float(row[3])})

rows = []
for row in dishes_ws.iter_rows(min_row=5, values_only=True):
    if not row[0]:
        continue
    name = str(row[3]).strip()
    rows.append(
        {
            "code": row[1],
            "category": row[2],
            "name": name,
            "ingredientsText": row[4],
            "servingGrams": float(row[6]),
            "kcal": float(row[7]),
            "proteinG": num(row[8]) or 0.0,
            "fatG": float(row[9]),
            "carbG": num(row[10]) or 0.0,
            "fiberG": num(row[11]),
            "sodiumMg": num(row[16]),
            "components": components.get(name, []),
        }
    )

missing = [r["name"] for r in rows if not r["components"]]
if missing:
    sys.exit(f"Món không có dòng chi tiết nguyên liệu: {missing[:5]}")

body = json.dumps(rows, ensure_ascii=False, indent=1)
OUT.write_text(
    "// FILE SINH TỰ ĐỘNG bởi scripts/import-vdd-xlsx.py từ packages/seed/source/vdd-tong-hop.xlsx.\n"
    "// ĐỪNG SỬA TAY. Gram từng nguyên liệu là số ƯỚC TÍNH; số dinh dưỡng của món là số bảng VDD.\n"
    "\n"
    "export interface VddRow {\n"
    "  code: string\n"
    "  category: string\n"
    "  name: string\n"
    "  ingredientsText: string\n"
    "  /** Khối lượng khẩu phần mà các số dinh dưỡng dưới đây ghi theo. */\n"
    "  servingGrams: number\n"
    "  kcal: number\n"
    "  proteinG: number\n"
    "  fatG: number\n"
    "  carbG: number\n"
    "  fiberG: number | null\n"
    "  sodiumMg: number | null\n"
    "  components: { name: string; grams: number }[]\n"
    "}\n\n"
    f"export const VDD_ROWS: readonly VddRow[] = {body}\n",
    encoding="utf-8",
)
print(f"Đã ghi {OUT} ({len(rows)} dòng, {sum(len(r['components']) for r in rows)} thành phần)")
